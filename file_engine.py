import csv
import io
import json
import os
import zipfile
from pathlib import Path
from typing import Any

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


OUTPUT_DIR = Path(
    os.getenv(
        "FILE_OUTPUT_DIR",
        "/tmp/quantum_queen_files",
    )
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "txt",
    "md",
    "rtf",
    "xlsx",
    "csv",
    "pptx",
    "html",
    "css",
    "js",
    "py",
    "json",
    "xml",
    "yaml",
    "yml",
    "sql",
    "zip",
}


def clean_filename(
    filename: str,
) -> str:

    filename = filename.strip()

    filename = filename.replace(
        "\\",
        "_",
    ).replace(
        "/",
        "_",
    )

    return filename or "quantum_queen_file"


def validate_extension(
    extension: str,
) -> str:

    extension = (
        extension
        .lower()
        .strip()
        .lstrip(".")
    )

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: .{extension}"
        )

    return extension


def create_text_file(
    content: str,
    filename: str,
    extension: str,
) -> Path:

    extension = validate_extension(
        extension
    )

    filename = clean_filename(
        filename
    )

    path = (
        OUTPUT_DIR
        / f"{filename}.{extension}"
    )

    path.write_text(
        content,
        encoding="utf-8",
    )

    return path


def create_pdf(
    content: str,
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.pdf"
    )

    pdf = canvas.Canvas(
        str(path),
        pagesize=A4,
    )

    width, height = A4

    x = 50
    y = height - 50

    for line in content.splitlines():

        if y < 50:
            pdf.showPage()
            y = height - 50

        pdf.drawString(
            x,
            y,
            line[:110],
        )

        y -= 16

    pdf.save()

    return path


def create_docx(
    content: str,
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.docx"
    )

    document = Document()

    for paragraph in content.splitlines():
        document.add_paragraph(
            paragraph
        )

    document.save(
        str(path)
    )

    return path


def create_xlsx(
    rows: list[list[Any]],
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.xlsx"
    )

    workbook = Workbook()

    sheet = workbook.active

    for row in rows:
        sheet.append(row)

    workbook.save(
        str(path)
    )

    return path


def create_csv(
    rows: list[list[Any]],
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.csv"
    )

    with path.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:

        writer = csv.writer(file)

        writer.writerows(rows)

    return path


def create_pptx(
    slides: list[dict[str, str]],
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.pptx"
    )

    presentation = Presentation()

    for slide_data in slides:

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        slide.shapes.title.text = (
            slide_data.get(
                "title",
                "",
            )
        )

        slide.placeholders[1].text = (
            slide_data.get(
                "content",
                "",
            )
        )

    presentation.save(
        str(path)
    )

    return path


def create_json(
    data: Any,
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.json"
    )

    path.write_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    return path


def create_zip(
    files: list[Path],
    filename: str,
) -> Path:

    path = (
        OUTPUT_DIR
        / f"{clean_filename(filename)}.zip"
    )

    with zipfile.ZipFile(
        path,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as archive:

        for file_path in files:

            archive.write(
                file_path,
                arcname=file_path.name,
            )

    return path


def create_file(
    file_type: str,
    filename: str,
    content: Any,
) -> dict[str, Any]:

    extension = validate_extension(
        file_type
    )

    if extension == "pdf":

        path = create_pdf(
            str(content),
            filename,
        )

    elif extension == "docx":

        path = create_docx(
            str(content),
            filename,
        )

    elif extension == "xlsx":

        path = create_xlsx(
            content,
            filename,
        )

    elif extension == "csv":

        path = create_csv(
            content,
            filename,
        )

    elif extension == "pptx":

        path = create_pptx(
            content,
            filename,
        )

    elif extension == "json":

        path = create_json(
            content,
            filename,
        )

    else:

        path = create_text_file(
            str(content),
            filename,
            extension,
        )

    return {
        "status": "created",
        "filename": path.name,
        "path": str(path),
        "extension": extension,
        "size": path.stat().st_size,
    }